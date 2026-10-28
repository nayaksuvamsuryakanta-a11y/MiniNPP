from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


OUTPUT = Path(__file__).with_name("MiniNPP_GUI_Presentation.pptx")
NAVY = RGBColor(20, 35, 62)
BLUE = RGBColor(46, 116, 181)
LIGHT = RGBColor(242, 246, 250)
TEXT = RGBColor(35, 45, 58)
MUTED = RGBColor(95, 108, 123)


def set_background(slide, color=LIGHT):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_header(slide, title):
    bar = slide.shapes.add_shape(1, 0, 0, Inches(13.333), Inches(0.18))
    bar.fill.solid()
    bar.fill.fore_color.rgb = BLUE
    bar.line.fill.background()
    box = slide.shapes.add_textbox(Inches(0.7), Inches(0.55), Inches(12), Inches(0.7))
    p = box.text_frame.paragraphs[0]
    p.text = title
    p.font.name = "Aptos Display"
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = NAVY


def add_bullets(slide, items, top=1.6):
    box = slide.shapes.add_textbox(Inches(0.9), Inches(top), Inches(6.1), Inches(4.5))
    tf = box.text_frame
    tf.word_wrap = True
    for index, item in enumerate(items):
        p = tf.paragraphs[0] if index == 0 else tf.add_paragraph()
        p.text = item
        p.level = 0
        p.font.name = "Aptos"
        p.font.size = Pt(20)
        p.font.color.rgb = TEXT
        p.space_after = Pt(14)


def add_placeholder(slide, text):
    shape = slide.shapes.add_shape(5, Inches(7.25), Inches(1.7), Inches(5.2), Inches(3.55))
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(224, 232, 241)
    shape.line.color.rgb = BLUE
    tf = shape.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = text
    p.alignment = PP_ALIGN.CENTER
    p.font.name = "Aptos"
    p.font.size = Pt(18)
    p.font.italic = True
    p.font.color.rgb = MUTED


def add_notes(slide, text):
    notes = slide.notes_slide.notes_text_frame
    notes.clear()
    notes.text = text


def add_content_slide(prs, title, bullets, placeholder, notes=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide)
    add_header(slide, title)
    add_bullets(slide, bullets)
    add_placeholder(slide, placeholder)
    if notes:
        add_notes(slide, notes)
    return slide


def main():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide, NAVY)
    accent = slide.shapes.add_shape(1, Inches(0.8), Inches(1.0), Inches(0.16), Inches(5.2))
    accent.fill.solid()
    accent.fill.fore_color.rgb = BLUE
    accent.line.fill.background()
    title = slide.shapes.add_textbox(Inches(1.35), Inches(1.45), Inches(11), Inches(1.4))
    p = title.text_frame.paragraphs[0]
    p.text = "MiniNPP: A Native Win32 Text Editor"
    p.font.name = "Aptos Display"
    p.font.size = Pt(34)
    p.font.bold = True
    p.font.color.rgb = RGBColor(255, 255, 255)
    subtitle = slide.shapes.add_textbox(Inches(1.38), Inches(3.25), Inches(10.5), Inches(1.2))
    p = subtitle.text_frame.paragraphs[0]
    p.text = "Built in C99, Scintilla, and Lexilla.\nFrom Disaster to Deployment."
    p.font.name = "Aptos"
    p.font.size = Pt(22)
    p.font.color.rgb = RGBColor(204, 219, 235)
    add_notes(slide, "Welcome to the presentation on MiniNPP, a native Win32 text editor built in C99.")

    add_content_slide(
        prs,
        "The Graphical User Interface",
        ["Raw Win32 API Architecture", "Scintilla Editing Engine", "Lexilla Syntax Highlighting", "Docked Status Bar & Tabs"],
        "[Insert your main GUI screenshot here]",
    )
    add_content_slide(
        prs,
        "Editing & UI Features",
        ["Syntax Highlighting & Code Folding", "Multi-Cursor & Column Editing", "Auto-completion & Bracket Matching"],
        "[Insert C code folding screenshot here]",
    )
    add_content_slide(
        prs,
        "Power User Features",
        ["Tabs with Session Restore (Unsaved tabs survive crashes)", "Find/Replace with Regex", "Macro Recording & Playback"],
        "[Insert JSON plugin or macro screenshot here]",
    )
    add_content_slide(
        prs,
        "Under the Hood & Disaster Recovery",
        ["Strict C99 Standard", "Dynamically loaded DLL Plugins", "Survived a catastrophic workspace wipe", "Fixed Win32 namespace collisions & memory leaks"],
        "Recovery timeline and architecture visual",
        "Explain the Remove-Item disaster and how we recovered from chat history.",
    )

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide, NAVY)
    title = slide.shapes.add_textbox(Inches(1), Inches(1.75), Inches(11.5), Inches(1.2))
    p = title.text_frame.paragraphs[0]
    p.text = "Thank You!"
    p.alignment = PP_ALIGN.CENTER
    p.font.name = "Aptos Display"
    p.font.size = Pt(40)
    p.font.bold = True
    p.font.color.rgb = RGBColor(255, 255, 255)
    subtitle = slide.shapes.add_textbox(Inches(1), Inches(3.25), Inches(11.5), Inches(1))
    p = subtitle.text_frame.paragraphs[0]
    p.text = "Check out the code on GitHub: github.com/nayaksuvamsuryakanta-a11y/MiniNPP"
    p.alignment = PP_ALIGN.CENTER
    p.font.name = "Aptos"
    p.font.size = Pt(20)
    p.font.color.rgb = RGBColor(204, 219, 235)

    prs.save(OUTPUT)
    print(f"Generated {OUTPUT}")


if __name__ == "__main__":
    main()
